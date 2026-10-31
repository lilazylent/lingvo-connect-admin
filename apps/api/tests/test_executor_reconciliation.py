"""Executor reconciliation export (paid package 2026-10-31): Oleg's underlined columns."""

from io import BytesIO

from openpyxl import load_workbook

from app.crm_stats import RECONCILIATION_HEADER
from tests.test_crm import setup

PATH = "/api/admin/crm/statistics/executor-reconciliation.xlsx"


def _rows(response):
    assert response.status_code == 200, response.text
    sheet = load_workbook(BytesIO(response.content)).worksheets[0]
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def test_reconciliation_has_one_row_per_executor_slice(client, create_user):
    headers = setup(client, create_user)
    assert client.post("/api/admin/crm/services", headers=headers, json={"code": "written_translation", "name": "Письменный перевод", "billing_mode": "CONDITIONAL_PAGE"}).status_code == 201
    ivanov = client.post("/api/admin/executors", headers=headers, json={"name": "Иванов"}).json()
    petrov = client.post("/api/admin/executors", headers=headers, json={"name": "Петров"}).json()
    assignment = {"billing_unit": "CONDITIONAL_PAGE"}
    order = client.post("/api/admin/crm/orders/wizard", headers=headers, json={"number": "26-0996", "works": [
        {"service_code": "written_translation", "source_language": "Английский", "target_language": "Русский", "topic": "техническая",
         "character_count": 30991, "client_rate": "590", "document_names": "601383 MultiPep 1 Manual Rev 1",
         "executor_assignments": [
             {**assignment, "executor_id": ivanov["id"], "character_count": 20000, "rate": "230", "deadline": "2026-01-12", "notes": "первая часть"},
             {**assignment, "executor_id": petrov["id"], "character_count": 10991, "rate": "150"},
             {**assignment, "executor_id": petrov["id"], "character_count": 500, "rate": "150", "status": "CANCELLED"},
         ]},
    ]})
    assert order.status_code == 201, order.text

    rows = _rows(client.get(PATH))
    assert rows[0] == RECONCILIATION_HEADER
    body = rows[1:]
    # the cancelled slice is not part of the reconciliation
    assert [(r[0], r[5]) for r in body] == [("26-0996", "Иванов"), ("26-0996", "Петров")]
    first = body[0]
    assert first[2:5] == ["Английский", "Русский", "техническая"]
    assert first[7] == "первая часть" and first[8] == 230 and first[9] == 20000
    # executor pages without the client minimum; payable = rate × pages (calculated, not «paid»)
    assert first[10] == 11.2 and first[11] == 2576 and first[12] == "601383 MultiPep 1 Manual Rev 1"

    only_petrov = _rows(client.get(f"{PATH}?executor_id={petrov['id']}"))[1:]
    assert [r[5] for r in only_petrov] == ["Петров"] and only_petrov[0][11] == 930


def test_reconciliation_skips_works_without_executors(client, create_user):
    headers = setup(client, create_user)
    order = client.post("/api/admin/crm/orders/wizard", headers=headers, json={"works": [{"work_type": "custom", "billing_unit": "FIXED", "client_rate": "100"}]})
    assert order.status_code == 201
    assert _rows(client.get(PATH))[1:] == []
