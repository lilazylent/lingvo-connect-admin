"""Owner feedback: editable number/date, client ranking, statistics and Excel exports."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from tests.test_crm import setup


def _order(client, headers, *, price, status="NEW", client_id=None, executor_id=None, cost=None, number=None, order_date=None, service=""):
    work = {"price": str(price), "service_code": service, "page_count": "2", "topic": "юридическая"}
    if executor_id:
        work["executor_assignments"] = [{"executor_id": executor_id, "cost": str(cost if cost is not None else price * 0.4)}]
    payload = {"client_id": client_id, "works": [work]}
    if number:
        payload["number"] = number
    if order_date:
        payload["order_date"] = order_date
    created = client.post("/api/admin/crm/orders/wizard", headers=headers, json=payload)
    assert created.status_code == 201, created.text
    order = created.json()
    if status != "NEW":
        assert client.patch(f"/api/admin/crm/orders/{order['id']}/status", headers=headers, json={"status": status}).status_code == 200
    return order


def test_order_number_and_date_are_editable_and_counter_stays_ahead(client, create_user):
    headers = setup(client, create_user)
    manual = _order(client, headers, price=100, number="26-0500", order_date="2026-03-15")
    assert manual["number"] == "26-0500"
    assert manual["created_at"].startswith("2026-03-15")
    # the next automatic number continues after the manual one
    assert client.get("/api/admin/crm/orders/number-preview?execution_year=2026").json()["number"] == "26-0501"
    # duplicates and bad formats are rejected
    dup = client.post("/api/admin/crm/orders/wizard", headers=headers, json={"number": "26-0500"})
    assert dup.status_code == 409
    assert client.post("/api/admin/crm/orders/wizard", headers=headers, json={"number": "26-5"}).status_code == 422
    # editing the number and the date of an existing order
    other = _order(client, headers, price=100)
    edited = client.patch(f"/api/admin/crm/orders/{other['id']}", headers=headers, json={"client_id": None, "contact_id": None, "manager_id": None, "notes": "", "number": "26-0700", "order_date": "2026-02-01", "version": other["version"]})
    assert edited.status_code == 200, edited.text
    assert edited.json()["number"] == "26-0700" and edited.json()["created_at"].startswith("2026-02-01")
    taken = client.patch(f"/api/admin/crm/orders/{other['id']}", headers=headers, json={"client_id": None, "contact_id": None, "manager_id": None, "notes": "", "number": "26-0500", "version": edited.json()["version"]})
    assert taken.status_code == 409
    assert client.get("/api/admin/crm/orders/number-preview?execution_year=2026").json()["number"] == "26-0701"


def test_clients_rank_by_metrics_and_filter_by_kind(client, create_user):
    headers = setup(client, create_user)
    big = client.post("/api/admin/companies", headers=headers, json={"name": "Альфа", "kind": "company"}).json()
    small = client.post("/api/admin/companies", headers=headers, json={"name": "Бета", "kind": "company"}).json()
    person = client.post("/api/admin/companies", headers=headers, json={"name": "Иванов Иван", "kind": "individual"}).json()
    executor = client.post("/api/admin/executors", headers=headers, json={"name": "Переводчик"}).json()
    _order(client, headers, price=1000, status="COMPLETED", client_id=small["id"], executor_id=executor["id"], cost=100)
    for _ in range(2):
        _order(client, headers, price=3000, status="COMPLETED", client_id=big["id"], executor_id=executor["id"], cost=2500)
    _order(client, headers, price=9000, status="ESTIMATING", client_id=person["id"])  # not accepted -> no revenue

    by_revenue = client.get("/api/admin/companies?sort=revenue&order=desc").json()["items"]
    assert [c["name"] for c in by_revenue] == ["Альфа", "Бета", "Иванов Иван"]
    assert Decimal(str(by_revenue[0]["stats"]["revenue"])) == Decimal("6000")
    by_profit = client.get("/api/admin/companies?sort=profit&order=desc").json()["items"]
    assert [c["name"] for c in by_profit][:2] == ["Альфа", "Бета"] and Decimal(str(by_profit[1]["stats"]["profit"])) == Decimal("900")
    by_orders = client.get("/api/admin/companies?sort=orders&order=desc").json()["items"]
    assert by_orders[0]["name"] == "Альфа" and by_orders[0]["stats"]["orders_count"] == 2
    individuals = client.get("/api/admin/companies?kind=individual").json()
    assert [c["name"] for c in individuals["items"]] == ["Иванов Иван"] and individuals["total"] == 1
    summary = client.get(f"/api/admin/crm/clients/{big['id']}/summary").json()
    assert Decimal(str(summary["profit"])) == Decimal("1000")


def test_statistics_and_excel_exports(client, create_user):
    headers = setup(client, create_user)
    company = client.post("/api/admin/companies", headers=headers, json={"name": "Клиент статистики"}).json()
    executor = client.post("/api/admin/executors", headers=headers, json={"name": "Иванов"}).json()
    _order(client, headers, price=1000, status="COMPLETED", client_id=company["id"], executor_id=executor["id"], cost=400, order_date="2026-01-10")
    _order(client, headers, price=2000, status="IN_PROGRESS", client_id=company["id"], order_date="2026-02-10")
    _order(client, headers, price=5000, status="ESTIMATING", client_id=company["id"], order_date="2026-02-11")
    _order(client, headers, price=700, status="CANCELLED", client_id=company["id"], order_date="2026-02-12")

    stats = client.get("/api/admin/crm/statistics").json()
    totals = stats["totals"]
    assert totals["orders"] == 2 and Decimal(str(totals["revenue"])) == Decimal("3000")
    assert Decimal(str(totals["executor_cost"])) == Decimal("400") and Decimal(str(totals["profit"])) == Decimal("2600")
    assert totals["pages"] == 4 and stats["by_topic"][0]["topic"] == "юридическая"
    assert [m["month"] for m in stats["by_month"]] == ["2026-01", "2026-02"]
    february = client.get("/api/admin/crm/statistics?date_from=2026-02-01&date_to=2026-02-28").json()["totals"]
    assert Decimal(str(february["revenue"])) == Decimal("2000")
    by_executor = client.get(f"/api/admin/crm/statistics?executor_id={executor['id']}").json()
    assert Decimal(str(by_executor["totals"]["executor_cost"])) == Decimal("400") and by_executor["by_executor"][0]["name"] == "Иванов"

    export = client.get("/api/admin/crm/statistics/export.xlsx?date_from=2026-01-01&date_to=2026-12-31")
    assert export.status_code == 200 and "spreadsheetml" in export.headers["content-type"]
    wb = load_workbook(BytesIO(export.content))
    assert {"Заказы", "Сводка", "По услугам", "По клиентам", "По исполнителям"} <= set(wb.sheetnames)
    assert wb["Заказы"].max_row == 1 + 3  # header + three non-cancelled orders (one work each)

    orders_export = client.get("/api/admin/crm/orders/export.xlsx?status=COMPLETED")
    assert orders_export.status_code == 200
    rows = list(load_workbook(BytesIO(orders_export.content))["Заказы"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1 and rows[0][2] == "Завершён" and rows[0][13] == "Иванов"
    assert date.fromisoformat(str(rows[0][1])[:10]) == date(2026, 1, 10)
