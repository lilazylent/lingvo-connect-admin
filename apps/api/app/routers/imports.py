"""Bounded spreadsheet import: preview first, explicit admin confirmation second."""

import io
import json
import re
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from typing import Annotated, Literal

import openpyxl
import xlrd
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, ValidationError
from sqlalchemy import func, select, text, update
from sqlalchemy.exc import IntegrityError

from app.client_models import ClientActivity, Company, Representative
from app.dependencies import SessionContext, as_utc, require_admin, require_admin_write
from app.import_models import ImportBatch
from app.models import utcnow
from app.operations_models import Executor, OperationalActivity, Order
from app.order_numbering import next_order_number
from app.routers.clients import DB, CompanyFields, PersonFields
from app.routers.operations import ExecutorFields, OrderFields, validate_order

router = APIRouter(prefix="/api/admin/imports", tags=["imports"])
Admin = Annotated[SessionContext, Depends(require_admin)]
AdminWrite = Annotated[SessionContext, Depends(require_admin_write)]
Entity = Literal["clients", "contacts", "executors", "orders"]
ALIASES = {
    "name": ["имя", "фио", "название", "название компании", "клиент", "исполнитель"],
    "email": ["email", "e-mail", "почта"],
    "phone": ["телефон"],
    "notes": ["комментарий", "примечание"],
    "tax_id": ["инн"],
    "kind": ["тип клиента"],
    "position": ["должность"],
    "telegram": ["телеграм"],
    "company": ["компания", "организация"],
    "client_id": ["id клиента"],
    "title": ["название заказа", "заказ"],
    "deadline": ["срок", "дедлайн"],
    "manager_id": ["id менеджера"],
    "status": ["статус"],
}


def cell(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return "" if value is None else str(value).strip()


def parse_table(content, filename):
    suffix = Path(filename).suffix.lower()
    rows = []
    try:
        if suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError("Распакованная таблица слишком велика (лимит 30 МБ)")
                if len(archive.infolist()) > 2000:
                    raise ValueError("Слишком сложная структура таблицы")
            book = openpyxl.load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            try:
                sheet = book.worksheets[0]
                if sheet.max_row and sheet.max_row > 501:
                    raise ValueError("За один импорт поддерживается до 500 строк")
                for row in sheet.iter_rows(max_row=501, max_col=41):
                    values = []
                    for item in row:
                        if item.data_type == "f":
                            raise ValueError(
                                "Таблица содержит формулы. Сохраните значения вместо формул."
                            )
                        values.append(cell(item.value))
                    rows.append(values)
            finally:
                book.close()
        elif suffix == ".xls":
            book = xlrd.open_workbook(file_contents=content, on_demand=True)
            try:
                sheet = book.sheet_by_index(0)
                if sheet.nrows > 501 or sheet.ncols > 40:
                    raise ValueError("Лимит таблицы: 500 строк и 40 столбцов")
                for i in range(sheet.nrows):
                    rows.append(
                        [
                            cell(xlrd.xldate.xldate_as_datetime(c.value, book.datemode))
                            if c.ctype == xlrd.XL_CELL_DATE
                            else cell(c.value)
                            for c in sheet.row(i)
                        ]
                    )
            finally:
                book.release_resources()
        else:
            raise ValueError("Выберите файл XLS или XLSX")
    except (ValueError, zipfile.BadZipFile, xlrd.XLRDError, KeyError, OSError) as exc:
        raise HTTPException(422, f"Не удалось прочитать таблицу: {exc}") from exc
    if not rows or not any(rows[0]):
        raise HTTPException(422, "Первая строка должна содержать названия столбцов")
    headers = rows[0]
    while headers and not headers[-1]:
        headers.pop()
    if len(headers) > 40 or len(set(headers)) != len(headers) or "" in headers:
        raise HTTPException(422, "Нужны уникальные непустые заголовки (до 40 столбцов)")
    return headers, rows[1:]


def resolve_company(db, values):
    company_name = values.pop("company", "")
    client_id = values.get("client_id")
    if not client_id:
        matches = list(
            db.scalars(
                select(Company)
                .where(Company.name == company_name, Company.archived.is_(False))
                .limit(2)
            )
        )
        if len(matches) != 1:
            raise ValueError("Компания не найдена или название неоднозначно. Укажите ID клиента.")
        client_id = matches[0].id
    company = db.get(Company, client_id)
    if not company or company.archived:
        raise ValueError("Клиент не найден или находится в архиве")
    values["client_id"] = client_id


def validate_row(db, entity, values, manager_id):
    if values.get("phone") and not re.fullmatch(r"\+?[\d\s()\-]{7,40}", values["phone"]):
        raise ValueError("Некорректный телефон")
    if entity in {"contacts", "orders"}:
        resolve_company(db, values)
    if entity == "clients":
        kinds = {
            "компания": "company",
            "физическое лицо": "individual",
            "частное лицо": "individual",
        }
        if "kind" in values:
            values["kind"] = kinds.get(values["kind"].lower(), values["kind"])
        model = CompanyFields.model_validate(values)
    elif entity == "contacts":
        client_id = values.pop("client_id")
        return {
            **PersonFields.model_validate(values).model_dump(mode="json"),
            "client_id": client_id,
        }
    elif entity == "executors":
        model = ExecutorFields.model_validate(values)
    else:
        values.setdefault("manager_id", manager_id)
        labels = {
            "новая": "NEW",
            "новый": "NEW",
            "в работе": "IN_PROGRESS",
            "завершён": "COMPLETED",
            "завершен": "COMPLETED",
            "отменён": "CANCELLED",
        }
        if "status" in values:
            values["status"] = labels.get(values["status"].lower(), values["status"])
        model = OrderFields.model_validate(values)
        validate_order(db, model)
    return model.model_dump(mode="json")


def duplicate(db, entity, row):
    if entity == "clients":
        filters = [Company.name == row["name"]]
        if row.get("tax_id"):
            filters = [Company.tax_id == row["tax_id"]]
        return db.scalar(select(Company.id).where(*filters).limit(1)) is not None
    if entity == "contacts":
        return (
            db.scalar(
                select(Representative.id)
                .where(
                    Representative.company_id == row["client_id"],
                    Representative.name == row["name"],
                    Representative.email == row["email"],
                )
                .limit(1)
            )
            is not None
        )
    if entity == "executors":
        return (
            db.scalar(
                select(Executor.id)
                .where(
                    Executor.name == row["name"], func.lower(Executor.email) == row["email"].lower()
                )
                .limit(1)
            )
            is not None
        )
    return (
        db.scalar(
            select(Order.id)
            .where(Order.client_id == row["client_id"], Order.title == row["title"])
            .limit(1)
        )
        is not None
    )


def batch_view(batch):
    return {
        "id": batch.id,
        "entity": batch.entity,
        "filename": batch.filename,
        "rows": batch.rows,
        "errors": batch.errors,
        "valid": len(batch.rows),
        "invalid": len(batch.errors),
        "total": len(batch.rows) + len(batch.errors),
        "report": batch.report,
    }


@router.post("/preview", status_code=201)
def preview(
    db: DB,
    context: AdminWrite,
    entity: Annotated[Entity, Form()],
    upload: Annotated[UploadFile, File()],
    mapping: Annotated[str, Form()] = "{}",
):
    content = upload.file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "Максимальный размер таблицы — 5 МБ")
    try:
        custom = json.loads(mapping)
        if not isinstance(custom, dict) or any(
            not isinstance(k, str) or not isinstance(v, str) for k, v in custom.items()
        ):
            raise ValueError()
    except (ValueError, TypeError) as exc:
        raise HTTPException(422, "Некорректное сопоставление столбцов") from exc
    headers, source = parse_table(content, upload.filename or "")
    aliases = {alias: key for key, values in ALIASES.items() for alias in [key, *values]}
    columns = [custom.get(h, aliases.get(h.lower(), "")) for h in headers]
    if any(v not in ALIASES and v for v in columns):
        raise HTTPException(422, "Сопоставление содержит неизвестное поле CRM")
    mapped = [c for c in columns if c]
    if len(mapped) != len(set(mapped)):
        raise HTTPException(422, "Два столбца сопоставлены одному полю CRM")
    required = "title" if entity == "orders" else "name"
    if required not in columns:
        raise HTTPException(
            422,
            f"Не найден столбец {'Название заказа' if entity == 'orders' else 'Имя / название'}",
        )
    valid, errors, seen = [], [], set()
    for number, cells in enumerate(source, 2):
        if not any(cells):
            continue
        try:
            values = {
                key: value for key, value in zip(columns, cells, strict=False) if key and value
            }
            result = validate_row(db, entity, values, context.user.id)
            signature = json.dumps(result, sort_keys=True, ensure_ascii=False)
            if signature in seen or duplicate(db, entity, result):
                raise ValueError("Похожая запись уже существует; строка будет пропущена")
            seen.add(signature)
            valid.append({"line": number, "data": result})
        except ValidationError as exc:
            message = "; ".join(f"{e['loc'][0]}: {e['msg']}" for e in exc.errors())
            errors.append({"line": number, "message": message})
        except (ValueError, HTTPException) as exc:
            errors.append(
                {
                    "line": number,
                    "message": str(exc.detail if isinstance(exc, HTTPException) else exc),
                }
            )
    batch = ImportBatch(
        owner_id=context.user.id,
        entity=entity,
        filename=Path(upload.filename or "table").name[:255],
        rows=valid,
        errors=errors,
        expires_at=utcnow() + timedelta(hours=2),
    )
    db.add(batch)
    db.commit()
    return batch_view(batch)


class Confirm(BaseModel):
    valid_rows: int


@router.post("/{batch_id}/confirm")
def confirm(batch_id: str, payload: Confirm, db: DB, context: AdminWrite):
    try:
        return apply_batch(batch_id, payload, db, context)
    except (IntegrityError, ValueError) as exc:
        db.rollback()
        raise HTTPException(
            409,
            "Данные изменились после проверки. Импорт отменён целиком; проверьте таблицу заново.",
        ) from exc


def apply_batch(batch_id, payload, db, context):
    if db.bind.dialect.name == "postgresql":
        # Serialize confirmations from different batches, avoiding cross-import duplicates.
        db.execute(text("SELECT pg_advisory_xact_lock(7142026)"))
    batch = db.scalar(
        select(ImportBatch)
        .where(ImportBatch.id == batch_id, ImportBatch.owner_id == context.user.id)
        .with_for_update()
    )
    if not batch:
        raise HTTPException(404, "Предпросмотр не найден")
    if batch.report is not None:
        return batch_view(batch)
    if as_utc(batch.expires_at) < utcnow():
        raise HTTPException(409, "Предпросмотр устарел. Загрузите таблицу заново.")
    if payload.valid_rows != len(batch.rows):
        raise HTTPException(409, "Число подтверждённых строк не совпадает с предпросмотром")
    created, skipped = [], []
    for item in batch.rows:
        values = dict(item["data"])
        if duplicate(db, batch.entity, values):
            skipped.append({"line": item["line"], "message": "Запись уже существует"})
            continue
        if batch.entity == "clients":
            row = Company(**CompanyFields.model_validate(values).model_dump())
        elif batch.entity == "contacts":
            resolve_company(db, values)
            row = Representative(company_id=values.pop("client_id"), **values)
        elif batch.entity == "executors":
            values.pop("directions", None)
            values.pop("position", None)
            row = Executor(**values)
        else:
            model = OrderFields.model_validate(values)
            validate_order(db, model)
            number = next_order_number(db, execution_year=model.deadline.year if model.deadline else None)
            row = Order(number=number, **model.model_dump())
        db.add(row)
        db.flush()
        if batch.entity in {"clients", "contacts"}:
            db.add(
                ClientActivity(
                    company_id=row.id if batch.entity == "clients" else row.company_id,
                    representative_id=row.id if batch.entity == "contacts" else None,
                    actor_user_id=context.user.id,
                    action="import.created",
                )
            )
        else:
            db.add(
                OperationalActivity(
                    actor_user_id=context.user.id,
                    action="import.created",
                    executor_id=row.id if batch.entity == "executors" else None,
                    order_id=row.id if batch.entity == "orders" else None,
                )
            )
        created.append({"line": item["line"], "id": row.id})
    batch.report = {"created": created, "skipped": skipped, "count": len(created)}
    db.commit()
    return batch_view(batch)
