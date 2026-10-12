"""Business statistics shared by the dashboard, clients list, statistics page and Excel export.

Money rules match the dashboard: revenue and profit only count orders the client has accepted
(``REVENUE_ORDER_STATUSES``); ``Новый``, ``Рассчитан`` and ``Отменён`` never become revenue.
Executor cost per work = active executor assignments, or the legacy per-work cost when a work
has no assignments (same as order ``finance()``).
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.client_models import Company
from app.crm_models import ClientPayment, ServiceType
from app.operations_models import Executor, ExecutorAssignment, Order, OrderWork

REVENUE_ORDER_STATUSES = ("APPROVED", "IN_PROGRESS", "REVIEW", "READY", "DELIVERED", "COMPLETED")
STATUS_LABELS = {
    "NEW": "Новый", "ESTIMATING": "Рассчитан", "APPROVED": "Согласован", "IN_PROGRESS": "В работе",
    "REVIEW": "На проверке", "READY": "Готов", "DELIVERED": "Выдан", "COMPLETED": "Завершён", "CANCELLED": "Отменён",
}
PAYMENT_METHODS = {"cash": "Наличные", "cashless": "Безналичный", "deposit": "Депозит"}
ZERO = Decimal("0")


def money(value) -> Decimal:
    return Decimal(str(value or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


@dataclass
class WorkRow:
    work: OrderWork
    order: Order
    price: Decimal
    cost: Decimal
    executors: list[tuple[str, Decimal]] = field(default_factory=list)


@dataclass
class Dataset:
    orders: list[Order]
    works: list[WorkRow]
    clients: dict[str, Company]
    payments: dict[str, ClientPayment]
    services: dict[str, str]


def load(
    db: Session,
    *,
    orders: list[Order] | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    client_id: str = "",
    executor_id: str = "",
    service_code: str = "",
    kind: str = "",
    include_cancelled: bool = False,
) -> Dataset:
    """Orders + works in scope. ``orders`` may be given directly (e.g. the orders-list filters)."""
    if orders is None:
        query = select(Order)
        if not include_cancelled:
            query = query.where(Order.status != "CANCELLED")
        if client_id:
            query = query.where(Order.client_id == client_id)
        if kind in ("company", "individual"):
            query = query.where(Order.client_id.in_(select(Company.id).where(Company.kind == kind)))
        orders = list(db.scalars(query.order_by(Order.created_at, Order.number)).all())
        if date_from:
            orders = [o for o in orders if o.created_at.date() >= date_from]
        if date_to:
            orders = [o for o in orders if o.created_at.date() <= date_to]
    order_ids = [o.id for o in orders]
    by_id = {o.id: o for o in orders}
    works = list(db.scalars(select(OrderWork).where(OrderWork.order_id.in_(order_ids), OrderWork.archived.is_(False)).order_by(OrderWork.sort_order)).all()) if order_ids else []
    work_ids = [w.id for w in works]
    assignments = list(db.scalars(select(ExecutorAssignment).where(ExecutorAssignment.work_id.in_(work_ids), ExecutorAssignment.archived.is_(False)).order_by(ExecutorAssignment.sort_order)).all()) if work_ids else []
    executor_ids = {a.executor_id for a in assignments} | {w.executor_id for w in works if w.executor_id}
    executor_names = {e.id: e.name for e in db.scalars(select(Executor).where(Executor.id.in_(executor_ids))).all()} if executor_ids else {}
    by_work: dict[str, list[ExecutorAssignment]] = defaultdict(list)
    for a in assignments:
        by_work[a.work_id].append(a)

    rows: list[WorkRow] = []
    for w in works:
        if service_code and w.service_code != service_code:
            continue
        work_assignments = by_work.get(w.id, [])
        if work_assignments:
            executors = [(executor_names.get(a.executor_id, ""), money(a.cost)) for a in work_assignments if not executor_id or a.executor_id == executor_id]
            if executor_id and not executors:
                continue
            cost = money(sum((c for _, c in executors), ZERO))
        else:
            if executor_id and w.executor_id != executor_id:
                continue
            cost = money(w.executor_cost)
            executors = [(executor_names.get(w.executor_id, ""), cost)] if w.executor_id else []
        rows.append(WorkRow(w, by_id[w.order_id], money(w.price) if w.client_billable else ZERO, cost, executors))
    if service_code or executor_id:
        kept = {r.order.id for r in rows}
        orders = [o for o in orders if o.id in kept]
    client_ids = {o.client_id for o in orders if o.client_id}
    clients = {c.id: c for c in db.scalars(select(Company).where(Company.id.in_(client_ids))).all()} if client_ids else {}
    payments = {p.order_id: p for p in db.scalars(select(ClientPayment).where(ClientPayment.order_id.in_([o.id for o in orders]))).all()} if orders else {}
    services = {s.code: s.name for s in db.scalars(select(ServiceType)).all()}
    return Dataset(orders, rows, clients, payments, services)


def volume(work: OrderWork) -> tuple[Decimal, Decimal, Decimal]:
    """(pages, documents, hours) in the work's own billing unit."""
    unit = work.billing_unit
    pages = Decimal(str(work.page_count or 0)) if unit in ("CONDITIONAL_PAGE", "PER_PAGE", "CUSTOM") else ZERO
    docs = Decimal(work.document_count or 0) if unit == "PER_DOCUMENT" else ZERO
    hours = Decimal(str(work.hour_count or 0)) if unit == "HOURLY" else ZERO
    return pages, docs, hours


def _bucket():
    return {"orders": set(), "works": 0, "pages": ZERO, "documents": ZERO, "hours": ZERO, "revenue": ZERO, "executor_cost": ZERO}


def _finish(key_fields: dict, b: dict) -> dict:
    revenue, cost = money(b["revenue"]), money(b["executor_cost"])
    return {**key_fields, "orders": len(b["orders"]), "works": b["works"], "pages": float(b["pages"]), "documents": float(b["documents"]),
            "hours": float(b["hours"]), "revenue": revenue, "executor_cost": cost, "profit": money(revenue - cost)}


def summarize(ds: Dataset) -> dict:
    accepted = {o.id for o in ds.orders if o.status in REVENUE_ORDER_STATUSES}
    total = _bucket()
    by_service, by_topic, by_client, by_executor, by_month = (defaultdict(_bucket) for _ in range(5))
    for r in ds.works:
        if r.order.id not in accepted:
            continue
        pages, docs, hours = volume(r.work)
        month = r.order.created_at.strftime("%Y-%m")
        for b in (total, by_service[r.work.service_code or ""], by_topic[(r.work.topic or "").strip()], by_client[r.order.client_id or ""], by_month[month]):
            b["orders"].add(r.order.id); b["works"] += 1
            b["pages"] += pages; b["documents"] += docs; b["hours"] += hours
            b["revenue"] += r.price; b["executor_cost"] += r.cost
        for name, cost in r.executors:
            eb = by_executor[name or "Без имени"]
            eb["orders"].add(r.order.id); eb["works"] += 1; eb["pages"] += pages; eb["documents"] += docs; eb["hours"] += hours
            eb["executor_cost"] += cost
    for o in ds.orders:  # orders without works still count as orders
        if o.id in accepted:
            total["orders"].add(o.id); by_client[o.client_id or ""]["orders"].add(o.id); by_month[o.created_at.strftime("%Y-%m")]["orders"].add(o.id)
    paid = money(sum((money(ds.payments[o].amount_paid) for o in accepted if o in ds.payments), ZERO))
    revenue = money(total["revenue"])
    totals = _finish({}, total)
    totals.update(
        orders_all=len(ds.orders), orders_not_accepted=len([o for o in ds.orders if o.id not in accepted]),
        paid=paid, debt=money(max(ZERO, revenue - paid)),
        margin_percent=money(totals["profit"] / revenue * 100) if revenue else ZERO,
    )
    clients = []
    for cid, b in by_client.items():
        c = ds.clients.get(cid)
        clients.append(_finish({"client_id": cid, "name": c.name if c else "Без клиента", "kind": c.kind if c else ""}, b))
    return {
        "totals": totals,
        "by_service": sorted((_finish({"service_code": k, "name": ds.services.get(k, k or "Без услуги")}, b) for k, b in by_service.items()), key=lambda x: -x["revenue"]),
        "by_topic": sorted((_finish({"topic": k or "Без тематики"}, b) for k, b in by_topic.items()), key=lambda x: -x["works"]),
        "by_client": sorted(clients, key=lambda x: -x["revenue"]),
        "by_executor": sorted((_finish({"name": k}, b) for k, b in by_executor.items()), key=lambda x: -x["executor_cost"]),
        "by_month": sorted((_finish({"month": k}, b) for k, b in by_month.items()), key=lambda x: x["month"]),
    }


def client_aggregates(db: Session) -> dict[str, dict]:
    """Per client: orders (not cancelled), accepted revenue, executor cost and profit."""
    ds = load(db)
    accepted = {o.id for o in ds.orders if o.status in REVENUE_ORDER_STATUSES}
    out: dict[str, dict] = defaultdict(lambda: {"orders_count": 0, "revenue": ZERO, "executor_cost": ZERO})
    for o in ds.orders:
        if o.client_id:
            out[o.client_id]["orders_count"] += 1
    for r in ds.works:
        if r.order.id in accepted and r.order.client_id:
            out[r.order.client_id]["revenue"] += r.price
            out[r.order.client_id]["executor_cost"] += r.cost
    for v in out.values():
        v["revenue"], v["executor_cost"] = money(v["revenue"]), money(v["executor_cost"])
        v["profit"] = money(v["revenue"] - v["executor_cost"])
    return out


# ---------------------------------------------------------------- Excel export
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="7A2945")


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: list[int] | None = None, first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = HEAD, HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    for i, width in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def workbook(ds: Dataset, title: str) -> bytes:
    """Orders as one row per work (client side + executor side), plus the summary sheets."""
    wb = Workbook()
    rows = []
    seen_orders = set()
    works_by_order: dict[str, list[WorkRow]] = defaultdict(list)
    for r in ds.works:
        works_by_order[r.order.id].append(r)
    for o in ds.orders:
        client = ds.clients.get(o.client_id or "")
        pay = ds.payments.get(o.id)
        for r in works_by_order.get(o.id) or [None]:
            first = o.id not in seen_orders
            seen_orders.add(o.id)
            w = r.work if r else None
            pages, docs, hours = volume(w) if w else (ZERO, ZERO, ZERO)
            rows.append([
                o.number, o.created_at.date(), STATUS_LABELS.get(o.status, o.status),
                client.name if client else "", ("Частное лицо" if client and client.kind == "individual" else "Компания") if client else "",
                ds.services.get(w.service_code, w.service_code) if w else "",
                " → ".join(x for x in (w.source_language, w.target_language) if x) if w else "",
                w.topic if w else "", float(pages) or None, float(docs) or None, float(hours) or None,
                money(w.client_rate) if w else None, r.price if r else None,
                ", ".join(n for n, _ in r.executors if n) if r else "", r.cost if r else None,
                (r.price - r.cost) if r else None,
                money(pay.amount_due) if pay and first else None, money(pay.amount_paid) if pay and first else None,
                PAYMENT_METHODS.get(pay.payment_method, pay.payment_method) if pay and first else "",
                pay.invoice_number if pay and first else "",
            ])
    _sheet(wb, "Заказы", ["№ заказа", "Дата", "Статус", "Клиент", "Тип клиента", "Услуга", "Языки", "Тематика", "Страниц", "Документов", "Часов",
                         "Тариф клиенту, ₽", "Стоимость для клиента, ₽", "Исполнитель", "Выплата исполнителю, ₽", "Прибыль по работе, ₽",
                         "К оплате по заказу, ₽", "Оплачено, ₽", "Способ оплаты", "№ счёта"], rows,
           [10, 11, 13, 28, 14, 24, 22, 18, 9, 11, 8, 12, 14, 26, 14, 14, 14, 12, 14, 10], first=True)
    for row in wb["Заказы"].iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "DD.MM.YY"
    s = summarize(ds)
    t = s["totals"]
    _sheet(wb, "Сводка", ["Показатель", "Значение"], [
        ["Период / фильтр", title], ["Принятых заказов", t["orders"]], ["Всего заказов в выборке (без отменённых)", t["orders_all"]],
        ["Выручка (принятые заказы), ₽", t["revenue"]], ["Оплачено клиентами, ₽", t["paid"]], ["Долг клиентов, ₽", t["debt"]],
        ["Выплаты исполнителям, ₽", t["executor_cost"]], ["Прибыль, ₽", t["profit"]], ["Маржинальность, %", t["margin_percent"]],
        ["Страниц", t["pages"]], ["Документов", t["documents"]], ["Часов", t["hours"]],
    ], [44, 22])
    cols = ["Заказов", "Работ", "Страниц", "Документов", "Часов", "Выручка, ₽", "Выплаты исполнителям, ₽", "Прибыль, ₽"]
    pick = lambda x: [x["orders"], x["works"], x["pages"], x["documents"], x["hours"], x["revenue"], x["executor_cost"], x["profit"]]  # noqa: E731
    _sheet(wb, "По услугам", ["Услуга", *cols], [[x["name"], *pick(x)] for x in s["by_service"]], [30] + [14] * 8)
    _sheet(wb, "По клиентам", ["Клиент", "Тип", *cols], [[x["name"], "Частное лицо" if x["kind"] == "individual" else "Компания", *pick(x)] for x in s["by_client"]], [30, 14] + [14] * 8)
    _sheet(wb, "По исполнителям", ["Исполнитель", "Заказов", "Работ", "Страниц", "Документов", "Часов", "Выплаты, ₽"],
           [[x["name"], x["orders"], x["works"], x["pages"], x["documents"], x["hours"], x["executor_cost"]] for x in s["by_executor"]], [30] + [14] * 6)
    _sheet(wb, "По тематикам", ["Тематика", *cols], [[x["topic"], *pick(x)] for x in s["by_topic"]], [30] + [14] * 8)
    _sheet(wb, "По месяцам", ["Месяц", *cols], [[x["month"], *pick(x)] for x in s["by_month"]], [12] + [14] * 8)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
