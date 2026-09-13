"""Conservative document metrics extraction used by CRM order files."""

from __future__ import annotations

import re
from pathlib import Path


def _text_metrics(text: str) -> tuple[int, int]:
    normalized = text.replace("\r\n", "\n")
    return len(normalized), len(re.findall(r"\S+", normalized))


def analyze_document(path: Path, original_name: str) -> dict[str, int | str | None]:
    suffix = Path(original_name).suffix.lower()
    result: dict[str, int | str | None] = {
        "page_count": None,
        "character_count": None,
        "word_count": None,
        "analysis_status": "UNSUPPORTED",
        "analysis_note": "Не удалось определить автоматически",
    }
    try:
        if suffix == ".txt":
            raw = path.read_bytes()
            text = None
            for encoding in ("utf-8-sig", "utf-8", "cp1251"):
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                return result
            chars, words = _text_metrics(text)
            result.update(character_count=chars, word_count=words, analysis_status="OK", analysis_note="Текст определён автоматически; количество страниц вводится вручную.")
            return result

        if suffix == ".docx":
            from docx import Document

            document = Document(path)
            text = "\n".join(p.text for p in document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    text += "\n" + "\t".join(cell.text for cell in row.cells)
            chars, words = _text_metrics(text)
            result.update(character_count=chars, word_count=words, analysis_status="PARTIAL", analysis_note="Текст DOCX определён автоматически. Число страниц без рендера ненадёжно — укажите вручную.")
            return result

        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            result["page_count"] = len(reader.pages)
            texts: list[str] = []
            for page in reader.pages:
                extracted = page.extract_text() or ""
                if extracted.strip():
                    texts.append(extracted)
            if texts:
                chars, words = _text_metrics("\n".join(texts))
                result.update(character_count=chars, word_count=words, analysis_status="OK", analysis_note="Страницы и текстовый слой PDF определены автоматически.")
            else:
                result.update(analysis_status="PARTIAL", analysis_note="Количество страниц определено. Текстового слоя нет; OCR автоматически не запускался.")
            return result

        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            chunks: list[str] = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    for value in row:
                        if value is not None:
                            chunks.append(str(value))
            chars, words = _text_metrics("\n".join(chunks))
            result.update(character_count=chars, word_count=words, analysis_status="PARTIAL", analysis_note="Текст ячеек XLSX определён автоматически. Страницы зависят от параметров печати и вводятся вручную.")
            return result
    except Exception:
        result.update(analysis_status="ERROR", analysis_note="Файл сохранён, но автоматический анализ не удался. Введите объём вручную.")
    return result
